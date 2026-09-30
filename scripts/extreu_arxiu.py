#!/usr/bin/env python3
"""
Extreu les dades dels fulls del GENERADOR_BARRUF de l'arxiu de l'AJUSC.

Cada full que calcula una edició porta l'estat anterior (`DadesÚltimBarruf`),
l'estat que en surt (`Dades9Barruf`) i el campionat. El campionat hi és en
dos formats, segons l'època:

  - `RESULTATS TORNEIG` (des de l'edició 121): una fila per partida, amb el
    resultat. Es prenen els noms ja comprovats pel full.
  - Només `DadesTorneig` (edicions 71 a 120): per a cada jugador, els rivals
    de cada ronda i el total de victòries. El resultat de cada partida no hi
    és.

Escriu un JSON per full, que llegeix `scripts/genera_arxiu.ts`. Llegir els
xlsx és el que triga; així només es fa un cop.

Ús:
    python3 scripts/extreu_arxiu.py <carpeta de l'arxiu> <carpeta de sortida>
"""

from __future__ import annotations

import datetime
import hashlib
import json
import re
import sys
import unicodedata
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")


def neteja(valor) -> str:
    """Com `netejaNom()`: fora els caràcters invisibles i els espais sobrers."""
    if valor is None:
        return ""
    text = "".join(
        " " if unicodedata.category(c) == "Zs" else c
        for c in str(valor)
        if unicodedata.category(c) != "Cf"
    )
    return re.sub(r"\s+", " ", text).strip()


def numero(valor):
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return valor
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        return None


def dades(ws) -> list[dict]:
    """Les files d'una pestanya DadesÚltimBarruf / Dades9Barruf, en ordre."""
    files = []
    for r in ws.iter_rows(min_row=2, max_col=11, values_only=True):
        nom = neteja(r[0])
        if not nom:
            continue
        files.append({
            "nom": nom,
            "barruf": numero(r[1]),
            "estat": neteja(r[2]) or None,
            "club": neteja(r[3]) or None,
            "vtemp": numero(r[4]) or 0,
            "ptemp": numero(r[5]) or 0,
            "vt": numero(r[6]) or 0,
            "pt": numero(r[7]) or 0,
            "deb": r[8] == "deb",
            "darrera": neteja(r[9]) or None,
            "posicio": numero(r[10]) if len(r) > 10 and not isinstance(r[10], bool) else None,
        })
    return files


def totals(ws) -> dict:
    capcalera = next(ws.iter_rows(max_row=1, values_only=True))
    out = {}
    for etiqueta, clau in (("Total res individuals", "resultats"), ("Total campionats", "campionats")):
        if etiqueta in capcalera:
            out[clau] = numero(capcalera[capcalera.index(etiqueta) + 1])
    return out


def resultats(ws) -> tuple[dict, list[list[str]], str]:
    """Partides de RESULTATS TORNEIG: [ronda, jugador1, punts1, jugador2, punts2]."""
    files = list(ws.iter_rows(min_row=2, values_only=True))
    if not files:
        return {}, [], "buit"
    capcalera = [neteja(x) for x in files[0]]
    posicions = [i for i, x in enumerate(capcalera) if x == "Jugador 1"]
    if not posicions:
        return {}, [], "sense capçalera"

    def llegeix(j1: int) -> list[list]:
        ronda = j1 - 1 if j1 >= 1 and capcalera[j1 - 1] == "Ronda" else 0
        out = []
        for r in files[1:]:
            r = list(r) + [None] * 30
            a, b = neteja(r[j1]), neteja(r[j1 + 2])
            if a:
                out.append([numero(r[ronda]), a, numero(r[j1 + 1]), b or None, numero(r[j1 + 3])])
        return out

    # Com es va picar cada nom i com és al BARRUF (columnes G i H).
    alies = []
    if "Jugadors" in capcalera and "Nom BARRUF" in capcalera:
        g, h = capcalera.index("Jugadors"), capcalera.index("Nom BARRUF")
        for r in files[1:]:
            r = list(r) + [None] * 30
            picat, barruf = neteja(r[g]), neteja(r[h])
            if picat and barruf and picat != barruf:
                alies.append([picat, barruf])

    originals = llegeix(posicions[0])
    comprovats = llegeix(posicions[1]) if len(posicions) > 1 else []
    # Els comprovats surten d'una fórmula de Google que, exportada a xlsx, a
    # vegades només en conserva la primera cel·la. Per això es desen també els
    # originals, amb els noms corregits segons la taula de comprovació: el
    # generador tria i la verificació diu si l'encerta.
    correccio = dict(alies)
    corregits = [[r[0], correccio.get(r[1], r[1]), r[2], correccio.get(r[3], r[3]) if r[3] else None, r[4]] for r in originals]
    if comprovats and len(comprovats) >= 0.9 * len(originals):
        return {"comprovats": comprovats, "originals": corregits}, alies, "comprovats"
    return {"comprovats": comprovats, "originals": corregits}, alies, "originals"


def aparellaments(ws) -> list[dict]:
    """DadesTorneig: per jugador, les victòries i el rival de cada ronda."""
    out = []
    for r in ws.iter_rows(min_row=7, max_col=33, values_only=True):
        nom = neteja(r[0])
        if not nom:
            continue
        rivals = [neteja(x) or None for x in r[2:33]]
        while rivals and rivals[-1] is None:
            rivals.pop()
        out.append({"nom": nom, "victories": numero(r[1]) or 0, "rivals": rivals})
    return out


def metadades(wb, cami: Path) -> dict:
    dt = list(wb["DadesTorneig"].iter_rows(max_row=2, max_col=8, values_only=True)) if "DadesTorneig" in wb.sheetnames else []
    fila1 = list(dt[0]) + [None] * 8 if dt else [None] * 8
    fila2 = list(dt[1]) + [None] * 8 if len(dt) > 1 else [None] * 8
    info = {"edicio": None, "campionat": neteja(fila1[3]) or None, "data": None, "temporada": neteja(fila2[7]) or None}

    text = neteja(fila1[2])
    m = re.search(r"(\d{2,3})", text)
    if m:
        info["edicio"] = int(m.group(1))
    if isinstance(fila2[2], (datetime.date, datetime.datetime)):
        info["data"] = fila2[2].strftime("%Y-%m-%d")

    if "INICI" in wb.sheetnames:
        inici = {neteja(r[0]): r[1] for r in wb["INICI"].iter_rows(max_col=2, values_only=True) if r[0]}
        info["campionat"] = neteja(inici.get("TORNEIG")) or info["campionat"]
        info["temporada"] = neteja(inici.get("TEMPORADA")) or info["temporada"]
        data = inici.get("DATA")
        if isinstance(data, (datetime.date, datetime.datetime)):
            info["data"] = data.strftime("%Y-%m-%d")
        anterior = next((v for k, v in inici.items() if k.startswith("També has de triar")), None)
        if numero(anterior):
            info["edicio"] = int(numero(anterior)) + 1

    # El número bo és el de la carpeta de l'edició: la casella «Nom del BARRUF»
    # sovint no s'actualitzava (el full de la 104 diu 103, el de la 123, 222).
    carpeta = re.search(r"(?i)barruf\s*(\d{2,3})", cami.parent.name)
    if carpeta:
        info["edicio_carpeta"] = int(carpeta.group(1))
    if info["edicio"] is None:
        m = re.search(r"(?i)b(\d{2,3})\b", cami.name) or re.search(r"(?i)barruf\D{0,4}(\d{2,3})", cami.name)
        if m:
            info["edicio"] = int(m.group(1))
    if carpeta:
        info["edicio"] = int(carpeta.group(1))
    if not info["temporada"]:
        m = re.search(r"(20\d\d-\d\d)", str(cami))
        info["temporada"] = m.group(1) if m else None
    return info


def extreu(cami: Path) -> dict | None:
    wb = openpyxl.load_workbook(cami, data_only=True, read_only=True)
    if not {"DadesÚltimBarruf", "Dades9Barruf"} <= set(wb.sheetnames):
        return None
    info = metadades(wb, cami)
    sortida = {
        "fitxer": str(cami),
        **info,
        "anterior": dades(wb["DadesÚltimBarruf"]),
        "nova": dades(wb["Dades9Barruf"]),
        "totals": totals(wb["DadesÚltimBarruf"]),
    }
    if "RESULTATS TORNEIG" in wb.sheetnames:
        versions, alies, origen = resultats(wb["RESULTATS TORNEIG"])
        if versions and (versions["comprovats"] or versions["originals"]):
            sortida.update(partides=versions, alies=alies, origen=origen)
    if not sortida.get("partides") and "DadesTorneig" in wb.sheetnames:
        sortida.update(aparellaments=aparellaments(wb["DadesTorneig"]), origen="DadesTorneig")
    return sortida


def main():
    arrel, desti = Path(sys.argv[1]), Path(sys.argv[2])
    desti.mkdir(parents=True, exist_ok=True)
    for cami in sorted(arrel.rglob("*.xlsx")):
        clau = hashlib.sha1(str(cami.relative_to(arrel)).encode()).hexdigest()[:12]
        fitxer = desti / f"{clau}.json"
        if fitxer.exists():
            continue
        try:
            dades_full = extreu(cami)
        except Exception as error:  # un full trencat no ha d'aturar la resta
            print(f"  ERROR {cami.relative_to(arrel)}: {error}", file=sys.stderr, flush=True)
            continue
        if dades_full is None:
            continue
        dades_full["fitxer"] = str(cami.relative_to(arrel))
        fitxer.write_text(json.dumps(dades_full, ensure_ascii=False))
        p = dades_full.get("partides") or {}
        n = len(p.get(dades_full.get("origen"), [])) if p else len(dades_full.get("aparellaments") or [])
        print(f"  {dades_full['edicio']}: {dades_full['fitxer'][:70]} · {dades_full.get('origen')} {n}", flush=True)


if __name__ == "__main__":
    main()
