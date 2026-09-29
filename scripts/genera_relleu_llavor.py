#!/usr/bin/env python3
"""
Genera la migració que canvia la llavor del BARRUF per una edició més nova.

Llegeix un full de l'AJUSC d'un campionat ja calculat, que porta dues edicions:

  DadesÚltimBarruf   l'edició anterior (la que s'ha triat a INICI!B6)
  Dades9Barruf       l'edició que en surt, la que s'ha publicat en PDF

La nova surt com a llavor, i l'anterior es desa com a edició d'arxiu. Amb les
dues, el PDF de la llavor es pot reproduir sencer: la progressió, les victòries
i partides des de l'anterior i la variació de posició surten de comparar-les.

La migració que en surt serveix igual per a una base de dades nova (on la
llavor d'abans ja s'hi haurà aplicat) com per a una que ja en tingui una: els
jugadors es reconeixen pel nom normalitzat, i els nous s'hi afegeixen. Es
planta si ja s'ha computat cap campionat, perquè llavors canviar el punt de
partida canviaria tot el que ja s'ha publicat.

Ús:
    python3 scripts/genera_relleu_llavor.py "Barruf 209 - ....xlsx" \\
        > supabase/migrations/<data>_llavor_<n>.sql
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from genera_llavor import (  # noqa: E402
    COLUMNES,
    DARRERA_TEMPORADA,
    PRIMERA_TEMPORADA,
    cita,
    neteja,
    normalitza,
    temporada,
)

# Jugadors que el full ha rebatejat respecte de la llavor anterior: nom nou → nom
# antic. Sense això es donarien d'alta com si fossin algú altre i la mateixa
# persona quedaria partida en dues fitxes. El nom antic queda com a àlies.
RENOMENATS = {
    "Jordi Plana Español": "Jordi Plana Espanyol",
    "Lourdes Sola": "Lourdes Sola López",
}

# La data exacta de publicació no és al full. Es posa la de creació del full de
# la 209 i queda dit a la descripció.
DATA_LLAVOR = date(2026, 9, 18)
DATA_ANTERIOR = date(2026, 9, 1)


def jugadors_de(full, amb_posicio: bool) -> list[dict]:
    jugadors = []
    for ordre, fila in enumerate(full.iter_rows(min_row=2, values_only=True), start=1):
        nom = fila[COLUMNES["nom"] - 1]
        if not nom:
            continue

        def valor(clau):
            return fila[COLUMNES[clau] - 1]

        codi, cohort = temporada(valor("darrera_temporada"))
        jugadors.append(
            {
                "ordre": ordre,
                "nom": neteja(str(nom)),
                "barruf": valor("barruf"),
                "estat": valor("estat"),
                "club": neteja(str(valor("club"))) if valor("club") else None,
                "victories_temporada": valor("victories_temporada") or 0,
                "partides_temporada": valor("partides_temporada") or 0,
                "victories_totals": valor("victories_totals") or 0,
                "partides_totals": valor("partides_totals") or 0,
                "debutant": valor("debutant") == "deb",
                "darrera_temporada": codi,
                "cohort_llegat": cohort,
                "posicio": valor("posicio") if amb_posicio else None,
            }
        )
    return jugadors


def posicions(jugadors: list[dict]) -> None:
    """Posició dels actius: els empatats la comparteixen, com fa el full (RANK)."""
    actius = [j["barruf"] for j in jugadors if j["estat"] == "act"]
    for j in jugadors:
        if j["estat"] == "act":
            j["posicio"] = 1 + sum(1 for b in actius if b > j["barruf"])


def llegeix(cami: str):
    wb = openpyxl.load_workbook(cami, data_only=True, read_only=True)

    inici = {fila[0]: fila[1] for fila in wb["INICI"].iter_rows(max_col=2, values_only=True) if fila[0]}
    anterior_num = int(next(v for k, v in inici.items() if str(k).startswith("També has de triar")))
    temporada_codi = str(next(v for k, v in inici.items() if k == "TEMPORADA"))

    capcalera = next(wb["DadesÚltimBarruf"].iter_rows(max_row=1, values_only=True))
    def despres(etiqueta):
        return capcalera[capcalera.index(etiqueta) + 1]
    anterior_totals = (int(despres("Total res individuals")), int(despres("Total campionats")))

    # La capçalera del PDF de la nova edició: número, campionat i totals.
    text = [str(f[0] or "") for f in wb["9Barruf"].iter_rows(max_row=5, max_col=1, values_only=True)]
    numero = int(re.search(r"Edició número (\d+)", " ".join(text)).group(1))
    campionat = next(t for t in text if t.startswith("Campionat computat:")).split(":", 1)[1].strip()
    totals = re.search(r"([\d.]+) resultats individuals anotats i (\d+) campionats", " ".join(text))
    nova_totals = (int(totals.group(1).replace(".", "")), int(totals.group(2)))

    anterior = jugadors_de(wb["DadesÚltimBarruf"], amb_posicio=True)
    nova = jugadors_de(wb["Dades9Barruf"], amb_posicio=False)
    posicions(nova)

    if numero != anterior_num + 1:
        raise SystemExit(f"L'edició {numero} no és la següent de la {anterior_num}")
    if [j["nom"] for j in nova] != [j["nom"] for j in anterior][: len(nova)]:
        raise SystemExit("Les dues pestanyes no tenen els jugadors en el mateix ordre")

    return {
        "temporada": temporada_codi,
        "anterior": {"numero": anterior_num, "totals": anterior_totals, "jugadors": anterior},
        "nova": {"numero": numero, "totals": nova_totals, "campionat": campionat, "jugadors": nova},
    }


def valors_sql(jugador: dict) -> str:
    return (
        f"{jugador['barruf']}, {int(jugador['partides_totals'])}, {jugador['victories_totals']}, "
        f"{int(jugador['partides_temporada'])}, {jugador['victories_temporada']}, "
        f"'{jugador['estat']}', {cita(jugador['darrera_temporada'])}, "
        f"{'TRUE' if jugador['cohort_llegat'] else 'FALSE'}, "
        f"{'TRUE' if jugador['debutant'] else 'FALSE'}, "
        f"{jugador['posicio'] if jugador['posicio'] is not None else 'NULL'}"
    )


def genera(dades: dict) -> str:
    anterior, nova = dades["anterior"], dades["nova"]

    vists = {}
    for j in nova["jugadors"]:
        clau = normalitza(j["nom"])
        if clau in vists:
            raise SystemExit(f"Dos jugadors comparteixen nom normalitzat: {vists[clau]!r} i {j['nom']!r}")
        vists[clau] = j["nom"]

    l: list[str] = []
    escriu = l.append

    escriu(f"""-- =============================================================================
-- Relleu de la llavor del BARRUF: edició {nova['numero']}
-- =============================================================================
-- Generat per scripts/genera_relleu_llavor.py a partir del full de l'AJUSC que
-- calcula l'edició {nova['numero']}. NO l'editeu a mà: torneu a executar el generador.
--
-- La {nova['numero']} passa a ser la llavor, el punt zero de la cadena. La
-- {anterior['numero']} es desa com a edició d'arxiu perquè el PDF de la llavor pugui
-- mostrar què es va moure respecte de l'anterior.
--
-- Els jugadors es reconeixen pel nom normalitzat. Els que no hi eren s'afegeixen
-- amb el número següent, i tothom pren l'ordre i el club que té al full.
-- =============================================================================

BEGIN;

DO $$
BEGIN
    IF EXISTS (SELECT 1 FROM campionats WHERE computa_barruf) THEN
        RAISE EXCEPTION 'Ja hi ha campionats computats: canviar la llavor ara canviaria '
                        'tot el que s''ha publicat. Cal fer-ho a mà i amb molt de compte.';
    END IF;
END $$;

INSERT INTO temporades (codi, any_inici, data_inici, data_fi) VALUES""")
    escriu(",\n".join(
        f"    ('{a}-{str(a + 1)[2:]}', {a}, DATE '{a}-09-01', DATE '{a + 1}-08-31')"
        for a in range(PRIMERA_TEMPORADA, DARRERA_TEMPORADA + 1)
    ))
    escriu("ON CONFLICT (codi) DO NOTHING;\n")

    clubs = sorted({j["club"] for j in nova["jugadors"] + anterior["jugadors"] if j["club"]})
    escriu("INSERT INTO clubs (nom) VALUES")
    escriu(",\n".join(f"    ({cita(c)})" for c in clubs))
    escriu("ON CONFLICT (nom) DO NOTHING;\n")

    escriu("""-- -----------------------------------------------------------------------------
-- La llista del full, en l'ordre del full
-- -----------------------------------------------------------------------------
CREATE TEMP TABLE llista_full (
    ordre           INTEGER PRIMARY KEY,
    nom             TEXT NOT NULL,
    nom_norm        TEXT NOT NULL,
    -- Si el full l'ha rebatejat, com es deia a la llavor anterior.
    nom_antic_norm  TEXT,
    club            TEXT,
    jugador_id      UUID
) ON COMMIT DROP;
""")
    escriu("INSERT INTO llista_full (ordre, nom, nom_norm, nom_antic_norm, club) VALUES")
    escriu(",\n".join(
        f"    ({j['ordre']}, {cita(j['nom'])}, {cita(normalitza(j['nom']))}, "
        f"{cita(normalitza(RENOMENATS[j['nom']]) if j['nom'] in RENOMENATS else None)}, "
        f"{cita(j['club'])})"
        for j in nova["jugadors"]
    ))
    escriu(";\n")

    escriu(f"""DO $$
DECLARE
    r    RECORD;
    v_id UUID;
BEGIN
    FOR r IN SELECT * FROM llista_full ORDER BY ordre LOOP
        SELECT jugador_id INTO v_id FROM jugador_alies WHERE alies_norm = r.nom_norm;

        -- Rebatejat: és la fitxa que ja hi havia, amb el nom nou.
        IF v_id IS NULL AND r.nom_antic_norm IS NOT NULL THEN
            SELECT jugador_id INTO v_id FROM jugador_alies WHERE alies_norm = r.nom_antic_norm;
            IF v_id IS NOT NULL THEN
                UPDATE jugadors SET nom_complet = r.nom WHERE id = v_id;
                INSERT INTO jugador_alies (jugador_id, alies, alies_norm, origen)
                VALUES (v_id, r.nom, r.nom_norm, 'llavor barruf {nova['numero']}');
            END IF;
        END IF;

        IF v_id IS NULL THEN
            INSERT INTO jugadors (nom_complet) VALUES (r.nom) RETURNING id INTO v_id;
            INSERT INTO jugador_alies (jugador_id, alies, alies_norm, origen)
            VALUES (v_id, r.nom, r.nom_norm, 'llavor barruf {nova['numero']}');
        END IF;

        UPDATE jugadors SET
            ordre_llista = r.ordre,
            club_id = (SELECT id FROM clubs WHERE nom = r.club)
        WHERE id = v_id;

        UPDATE llista_full SET jugador_id = v_id WHERE ordre = r.ordre;
    END LOOP;
END $$;

SELECT setval('ordre_llista_seq', (SELECT max(ordre_llista) FROM jugadors));

-- -----------------------------------------------------------------------------
-- Les edicions: fora les d'abans, i l'anterior i la llavor nova
-- -----------------------------------------------------------------------------
DELETE FROM barruf_edicions;

INSERT INTO barruf_edicions (
    numero, data_publicacio, temporada_codi, es_llavor, descripcio,
    campionats_computats, resultats_acumulats, campionats_acumulats
) VALUES
    ({anterior['numero']}, DATE '{DATA_ANTERIOR}', '{dades['temporada']}', FALSE,
     'Edició d''arxiu, heretada del full de l''AJUSC. Només serveix per comparar-hi la llavor. Data aproximada.',
     NULL, {anterior['totals'][0]}, {anterior['totals'][1]}),
    ({nova['numero']}, DATE '{DATA_LLAVOR}', '{dades['temporada']}', TRUE,
     'Punt de partida heretat del full de l''AJUSC. Sense partides al darrere: no es pot recalcular. Data aproximada.',
     {cita(nova['campionat'])}, {nova['totals'][0]}, {nova['totals'][1]});
""")

    for edicio in (anterior, nova):
        escriu(f"-- Valors de l'edició {edicio['numero']} ({len(edicio['jugadors'])})")
        escriu("""INSERT INTO barruf_valors (
    edicio_id, jugador_id, barruf,
    partides_totals, victories_totals, partides_temporada, victories_temporada,
    estat, darrera_temporada, cohort_llegat, debutant, posicio
)
SELECT e.id, l.jugador_id, v.barruf,
       v.pt, v.vt, v.ptemp, v.vtemp, v.estat::estat_jugador, v.darrera, v.cohort, v.deb, v.pos
FROM (VALUES""")
        files = []
        for j in edicio["jugadors"]:
            ordre = j["ordre"] if edicio is nova else next(
                (n["ordre"] for n in nova["jugadors"] if n["nom"] == j["nom"]), None
            )
            if ordre is None:
                raise SystemExit(f"{j['nom']!r} surt a l'edició anterior però no a la nova")
            files.append(f"    ({ordre}, {valors_sql(j)})")
        escriu(",\n".join(files))
        escriu(f""") AS v(ordre, barruf, pt, vt, ptemp, vtemp, estat, darrera, cohort, deb, pos)
JOIN llista_full l ON l.ordre = v.ordre
CROSS JOIN (SELECT id FROM barruf_edicions WHERE numero = {edicio['numero']}) e;
""")

    escriu("COMMIT;")
    return "\n".join(l) + "\n"


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    sys.stdout.write(genera(llegeix(sys.argv[1])))
